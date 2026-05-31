from pathlib import Path

from openpyxl import load_workbook
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
from pptx.util import Inches, Pt


BASE = Path(r"D:\Codex\tradecare-agent")
PAPER = BASE / "paper"
DESKTOP = Path.home() / "Desktop"

INPUT_PPTS = [
    (
        PAPER / "TradeCare-Agent_10min_visual_capture_RAG_sources_comparison_KR_FIXED.pptx",
        PAPER / "TradeCare-Agent_10min_visual_capture_RAG_sources_comparison_KR_FIXED_20QA.pptx",
    ),
    (
        PAPER / "Intelligent_TradeCare_Agent_RAG_sources_comparison_KR_FIXED.pptx",
        PAPER / "Intelligent_TradeCare_Agent_RAG_sources_comparison_KR_FIXED_20QA.pptx",
    ),
]


def load_cases():
    answer_file = next(path for path in DESKTOP.glob("*.xlsx") if path.stat().st_size == 11979)
    workbook = load_workbook(str(answer_file), data_only=True)
    rows = [row[0] for row in workbook.active.iter_rows(values_only=True) if row and row[0]]
    cases = []
    for index in range(0, len(rows), 3):
        chunk = [str(item) for item in rows[index : index + 3]]
        if len(chunk) < 3:
            continue
        case_id = chunk[0].replace("ID", "").strip()
        question = chunk[1].split("：", 1)[-1].strip()
        answer_core = chunk[2].split("：", 1)[-1].strip()
        cases.append((case_id, question, answer_core))
    return cases[:20]


def add_textbox(slide, x, y, w, h, text, size=18, bold=False, color=(15, 35, 60), align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    if align:
        paragraph.alignment = align
    run = paragraph.runs[0]
    run.font.name = "Malgun Gothic"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(*color)
    return box


def set_cell(cell, text, size=8.2, bold=False, color=(31, 41, 55), fill=None, align=None):
    if fill:
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(*fill)
    cell.margin_left = Inches(0.04)
    cell.margin_right = Inches(0.04)
    cell.margin_top = Inches(0.03)
    cell.margin_bottom = Inches(0.03)
    frame = cell.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    if align:
        paragraph.alignment = align
    paragraph.font.name = "Malgun Gothic"
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = RGBColor(*color)


def add_case_slide(prs, cases, title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = RGBColor(247, 248, 250)

    add_textbox(slide, 0.55, 0.25, 12.2, 0.28, "TradeCare-Agent", size=11, bold=True, color=(75, 94, 115))
    add_textbox(slide, 0.55, 0.62, 12.2, 0.45, title, size=25, bold=True, color=(12, 31, 54))
    add_textbox(slide, 0.58, 1.08, 12.0, 0.28, subtitle, size=11.5, color=(80, 88, 99))

    rows = len(cases) + 1
    table_shape = slide.shapes.add_table(rows, 3, Inches(0.55), Inches(1.55), Inches(12.25), Inches(4.95))
    table = table_shape.table
    table.columns[0].width = Inches(0.45)
    table.columns[1].width = Inches(5.3)
    table.columns[2].width = Inches(6.5)

    headers = ["No.", "고객 문의 예시", "예상 답변 핵심 / Gold Standard"]
    for col, header in enumerate(headers):
        set_cell(table.cell(0, col), header, size=10, bold=True, color=(255, 255, 255), fill=(29, 78, 116), align=PP_ALIGN.CENTER)

    for row_index, (case_id, question, answer_core) in enumerate(cases, 1):
        fill = (255, 255, 255) if row_index % 2 else (239, 245, 250)
        set_cell(table.cell(row_index, 0), case_id, size=8.4, bold=True, fill=fill, align=PP_ALIGN.CENTER)
        set_cell(table.cell(row_index, 1), question, size=8.0, fill=fill)
        set_cell(table.cell(row_index, 2), answer_core, size=8.0, fill=fill)

    note = (
        "활용 목적: 같은 20개 문의에 대해 일반 GPT, GPT+RAG, TradeCare Agent의 답변 품질을 비교하고, "
        "필수 정보 포함 여부와 위험 표현 통제를 평가합니다."
    )
    add_textbox(slide, 0.65, 6.67, 11.9, 0.32, note, size=10.5, bold=True, color=(55, 65, 81), align=PP_ALIGN.CENTER)
    return slide


def move_last_slide(prs, index):
    slide_id_list = prs.slides._sldIdLst
    slide_ids = list(slide_id_list)
    last = slide_ids[-1]
    slide_id_list.remove(last)
    slide_id_list.insert(min(index, len(slide_ids)), last)


def main():
    cases = load_cases()
    for source, target in INPUT_PPTS:
        prs = Presentation(str(source))
        add_case_slide(
            prs,
            cases[:10],
            "고객 예상 질문 20개 및 답변 핵심 (1/2)",
            "무역 고객이 실제로 물어볼 수 있는 질문과 Agent가 포함해야 할 답변 핵심입니다.",
        )
        add_case_slide(
            prs,
            cases[10:],
            "고객 예상 질문 20개 및 답변 핵심 (2/2)",
            "통관, 관세, 원산지, 운송 책임 등 고위험 상담은 확인·전문가 검토 신호를 포함해야 합니다.",
        )
        # Place the two slides after the qualitative answer-comparison slide.
        move_last_slide(prs, 11)
        move_last_slide(prs, 12)
        prs.save(str(target))
        print(target)


if __name__ == "__main__":
    main()
