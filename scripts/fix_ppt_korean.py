from pathlib import Path

from pptx import Presentation


BASE = Path(r"D:\Codex\tradecare-agent\paper")
FILES = [
    (
        BASE / "TradeCare-Agent_10min_visual_capture_RAG_sources_comparison.pptx",
        BASE / "TradeCare-Agent_10min_visual_capture_RAG_sources_comparison_KR_FIXED.pptx",
    ),
    (
        BASE / "Intelligent_TradeCare_Agent_RAG_sources_comparison.pptx",
        BASE / "Intelligent_TradeCare_Agent_RAG_sources_comparison_KR_FIXED.pptx",
    ),
]

REPLACEMENTS = {
    "RAG Knowledge Base Sources": "RAG 지식 베이스 출처",
    "Documented Yoon Hang Trade knowledge is retrieved before answer generation.": "답변 생성 전에 문서화된 윤항무역 지식을 먼저 검색합니다.",
    "Source Files": "출처 파일",
    "Curated From": "수집/정리 기준",
    "Yoon Hang Trade service scope": "윤항무역 서비스 범위",
    "Website-facing service information": "홈페이지 서비스 안내 정보",
    "Operator-provided counseling rules": "운영자가 제공한 상담 규칙",
    "Required intake fields by task": "업무 유형별 필수 확인 정보",
    "Rules against final price, MOQ, refund, customs guarantees before confirmation": "확인 전 단가, MOQ, 환불, 통관 확정 금지 규칙",
    "Single LLM ?? ??": "단일 LLM 대비 강점",
    "Single LLM 대비 강점": "단일 LLM 대비 강점",
    "Company policy grounding": "회사 정책 기반 답변",
    "Task-specific missing-field requests": "업무별 누락 정보 요청",
    "Less hallucinated capability claims": "과장된 가능 여부 답변 감소",
    "Risk and Critic checks after retrieval": "검색 후 Risk/Critic 검토 수행",
    "Example retrieval path": "예시 검색 흐름",
    "OEM/ODM inquiry -> services_oem_odm.md + quotation_required_fields.md -> asks for photos, specs, quantity, logo/package, lead time and avoids fixed commitments.": "OEM/ODM 문의 -> services_oem_odm.md + quotation_required_fields.md 검색 -> 사진, 사양, 수량, 로고/패키지, 납기를 요청하고 확정 표현을 피함",
    "Architecture justification: Retrieval provides evidence, Risk removes unsafe promises, Critic checks whether the answer follows retrieved constraints.": "아키텍처 정당성: Retrieval은 근거를 제공하고, Risk는 위험한 확정 표현을 제거하며, Critic은 답변이 검색된 제약을 따르는지 검토합니다.",
    "Answer Quality Comparison": "고객 응대 답변 품질 비교",
    "Same OEM/sample-production question, compared by required-info coverage and risk control.": "같은 OEM/샘플 제작 문의를 필수 정보 포함률과 위험 표현 통제 기준으로 비교합니다.",
    "General GPT": "일반 GPT",
    "Single response": "단일 답변",
    "Broad OEM explanation": "일반적인 OEM 설명 제공",
    "May miss images, specs, patterns, quantity, logo/package, destination": "이미지, 사양, 패턴, 수량, 로고/패키지, 배송지를 빠뜨릴 수 있음",
    "May sound too confident before supplier confirmation": "공급처 확인 전에도 지나치게 확정적으로 들릴 수 있음",
    "Coverage score: 0.139": "필수 정보 포함 점수: 0.139",
    "GPT + RAG": "GPT + RAG",
    "Retrieval only": "검색 기반 답변",
    "Uses service documents": "서비스 문서 참고",
    "Asks for some missing fields": "일부 누락 정보 요청",
    "No separate Risk Agent": "별도 Risk Agent 없음",
    "No Critic reflection loop": "Critic Reflection 루프 없음",
    "Coverage score: 0.678": "필수 정보 포함 점수: 0.678",
    "TradeCare Agent": "TradeCare Agent",
    "Multi-Agent + RAG + Reflection": "Multi-Agent + RAG + Reflection",
    "Classifies OEM/ODM intent": "OEM/ODM 의도 분류",
    "Retrieves company rules": "회사 업무 규칙 검색",
    "Requests concrete intake checklist": "구체적인 접수 체크리스트 요청",
    "Avoids unsupported MOQ, sample cost, lead time, feasibility promises": "MOQ, 샘플비, 납기, 제작 가능 여부를 확인 전 확정하지 않음",
    "Coverage score: 0.967": "필수 정보 포함 점수: 0.967",
    "Conclusion: TradeCare Agent turns an incomplete customer question into an actionable intake checklist while grounding the answer and preventing unsupported promises.": "결론: TradeCare Agent는 불완전한 고객 문의를 실행 가능한 접수 체크리스트로 바꾸고, 근거 기반 답변과 확정 표현 방지를 동시에 수행합니다.",
}


def replace_text_frame(text_frame):
    for paragraph in text_frame.paragraphs:
        full_text = "".join(run.text for run in paragraph.runs) if paragraph.runs else paragraph.text
        new_text = full_text
        for old, new in REPLACEMENTS.items():
            new_text = new_text.replace(old, new)
        if new_text != full_text:
            if paragraph.runs:
                paragraph.runs[0].text = new_text
                for run in paragraph.runs[1:]:
                    run.text = ""
            else:
                paragraph.text = new_text
        for run in paragraph.runs:
            run.font.name = "Malgun Gothic"


def main():
    for src, out in FILES:
        prs = Presentation(str(src))
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text_frame") and shape.has_text_frame:
                    replace_text_frame(shape.text_frame)
        prs.save(str(out))
        print(out)


if __name__ == "__main__":
    main()
