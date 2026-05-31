from pathlib import Path

from openpyxl import load_workbook
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
from pptx.util import Inches, Pt


BASE = Path(r"D:\Codex\tradecare-agent")
PAPER = BASE / "paper"
DESKTOP = Path.home() / "Desktop"

INPUT_PPTS = [
    (
        PAPER / "TradeCare-Agent_10min_visual_capture_RAG_sources_comparison_KR_FIXED.pptx",
        PAPER / "TradeCare-Agent_10min_visual_capture_RAG_sources_comparison_KR_FIXED_20QA_1slide.pptx",
    ),
    (
        PAPER / "Intelligent_TradeCare_Agent_RAG_sources_comparison_KR_FIXED.pptx",
        PAPER / "Intelligent_TradeCare_Agent_RAG_sources_comparison_KR_FIXED_20QA_1slide.pptx",
    ),
]


def load_cases():
    answer_file = next(path for path in DESKTOP.glob("*.xlsx") if path.stat().st_size == 11979)
    workbook = load_workbook(str(answer_file), data_only=True)
    rows = [row[0] for row in workbook.active.iter_rows(values_only=True) if row and row[0]]
    cases = []
    for index in range(0, len(rows), 3):
        chunk = [str(item) for item in rows[index : index + 3]]
        if len(chunk) < 3:
            continue
        case_id = chunk[0].replace("ID", "").strip()
        question = chunk[1].split("：", 1)[-1].strip()
        answer_core = chunk[2].split("：", 1)[-1].strip()
        cases.append((case_id, question, answer_core))
    return cases[:20]


def shorten(text, limit):
    text = text.replace("Gold Standard", "").replace("‌", "").strip()
    return text if len(text) <= limit else text[: limit - 1].rstrip() + "…"


def add_textbox(slide, x, y, w, h, text, size=18, bold=False, color=(15, 35, 60), align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    if align:
        paragraph.alignment = align
    run = paragraph.runs[0]
    run.font.name = "Malgun Gothic"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(*color)
    return box


def set_cell(cell, text, size=5.4, bold=False, color=(31, 41, 55), fill=None, align=None):
    if fill:
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(*fill)
    cell.margin_left = Inches(0.025)
    cell.margin_right = Inches(0.025)
    cell.margin_top = Inches(0.01)
    cell.margin_bottom = Inches(0.01)
    frame = cell.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    if align:
        paragraph.alignment = align
    paragraph.font.name = "Malgun Gothic"
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = RGBColor(*color)


def add_case_slide(prs, cases):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = RGBColor(247, 248, 250)

    add_textbox(slide, 0.45, 0.18, 12.4, 0.25, "TradeCare-Agent", size=10, bold=True, color=(75, 94, 115))
    add_textbox(slide, 0.45, 0.48, 12.4, 0.38, "고객 예상 질문 20개 및 예상 답변 핵심", size=22, bold=True, color=(12, 31, 54))
    add_textbox(slide, 0.48, 0.88, 12.2, 0.25, "평가 벤치마크에 사용 가능한 무역 고객 문의와 Agent가 포함해야 할 답변 핵심 요약입니다.", size=10.2, color=(80, 88, 99))

    rows = len(cases) + 1
    table_shape = slide.shapes.add_table(rows, 3, Inches(0.38), Inches(1.24), Inches(12.6), Inches(5.95))
    table = table_shape.table
    table.columns[0].width = Inches(0.38)
    table.columns[1].width = Inches(6.05)
    table.columns[2].width = Inches(6.17)

    for col, header in enumerate(["No.", "고객 문의", "예상 답변 핵심"]):
        set_cell(table.cell(0, col), header, size=7.1, bold=True, color=(255, 255, 255), fill=(29, 78, 116), align=PP_ALIGN.CENTER)

    for row_index, (case_id, question, answer_core) in enumerate(cases, 1):
        fill = (255, 255, 255) if row_index % 2 else (239, 245, 250)
        set_cell(table.cell(row_index, 0), case_id, size=5.7, bold=True, fill=fill, align=PP_ALIGN.CENTER)
        set_cell(table.cell(row_index, 1), shorten(question, 45), size=5.25, fill=fill)
        set_cell(table.cell(row_index, 2), shorten(answer_core, 55), size=5.25, fill=fill)


def move_last_slide(prs, index):
    slide_id_list = prs.slides._sldIdLst
    slide_ids = list(slide_id_list)
    last = slide_ids[-1]
    slide_id_list.remove(last)
    slide_id_list.insert(min(index, len(slide_ids)), last)


def main():
    cases = load_cases()
    for source, target in INPUT_PPTS:
        prs = Presentation(str(source))
        add_case_slide(prs, cases)
        move_last_slide(prs, 11)
        prs.save(str(target))
        print(target)


if __name__ == "__main__":
    main()
