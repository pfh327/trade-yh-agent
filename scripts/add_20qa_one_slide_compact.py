from pathlib import Path

from openpyxl import load_workbook
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
from pptx.util import Inches, Pt


BASE = Path(r"D:\Codex\tradecare-agent")
PAPER = BASE / "paper"
DESKTOP = Path.home() / "Desktop"

INPUT_PPTS = [
    (
        PAPER / "TradeCare-Agent_10min_visual_capture_RAG_sources_comparison_KR_FIXED.pptx",
        PAPER / "TradeCare-Agent_10min_visual_capture_RAG_sources_comparison_KR_FIXED_20QA_compact.pptx",
    ),
    (
        PAPER / "Intelligent_TradeCare_Agent_RAG_sources_comparison_KR_FIXED.pptx",
        PAPER / "Intelligent_TradeCare_Agent_RAG_sources_comparison_KR_FIXED_20QA_compact.pptx",
    ),
]

COMPACT_QUESTIONS = {
    "1": "세관 도착 후 배송 지연",
    "2": "중국 무역 계약서 검토",
    "3": "EU 의류 수출 통관서류",
    "4": "해상운송 발송인 정보 수정",
    "5": "신고 금액 오류 수정",
    "6": "식품 중국 수입 검역",
    "7": "세관 화물 압류 원인",
    "8": "FOB 비용 부담",
    "9": "안전한 결제 방식",
    "10": "운송 중 해손 배상",
    "11": "전자제품 중국 수입 관세",
    "12": "원산지증명서 재발급",
    "13": "혼적 화물 통관 방식",
    "14": "CIF/CFR 보험 책임",
    "15": "통관 지연 시 경매 처분",
    "16": "중국 업체 수출입 자격 조회",
    "17": "항공/해상 운임 계산",
    "18": "수하인 화물 포기 비용",
    "19": "FTA 관세 우대 증빙",
    "20": "관세 연체료 감면",
}

COMPACT_ANSWERS = {
    "1": "지연 원인 확인, 세관 문의, 추가서류 확인",
    "2": "필수 조항, 중국 법규, 전문가 검토",
    "3": "통관서류, EU 요건, 작성 유의점",
    "4": "수정 가능 여부, 절차, 추가비용",
    "5": "수정 신청, 필요서류, 과태료 가능성",
    "6": "검역 등록, 제출서류, 검사 항목",
    "7": "압류 원인, 처리절차, 기관 문의",
    "8": "발송인/수하인 비용 구분, 분쟁 기준",
    "9": "결제방식 위험도, 수수료, 추천 방식",
    "10": "책임 주체, 청구서류, 보험 처리",
    "11": "HS Code, FTA 조건, 세금 종류",
    "12": "재발급 조건, 신청절차, 유효기간",
    "13": "개별/공동 신고 조건과 장단점",
    "14": "책임 범위, 보험 주체, 비용 차이",
    "15": "초과 통관 절차, 경매 기준, 사전조치",
    "16": "공식 조회 채널, 절차, 핵심 확인정보",
    "17": "항공 부피중량, 해상 계산, 운임 기준",
    "18": "법적 책임, 부담 비용, 대응 방안",
    "19": "증빙자료, 원산지증명서, 신청 절차",
    "20": "감면 조건, 필요서류, 불가 사유",
}


def load_case_ids():
    answer_file = next(path for path in DESKTOP.glob("*.xlsx") if path.stat().st_size == 11979)
    workbook = load_workbook(str(answer_file), data_only=True)
    rows = [row[0] for row in workbook.active.iter_rows(values_only=True) if row and row[0]]
    case_ids = []
    for index in range(0, len(rows), 3):
        chunk = [str(item) for item in rows[index : index + 3]]
        if len(chunk) >= 3:
            case_ids.append(chunk[0].replace("ID", "").strip())
    return case_ids[:20]


def add_textbox(slide, x, y, w, h, text, size=18, bold=False, color=(15, 35, 60), align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    if align:
        paragraph.alignment = align
    run = paragraph.runs[0]
    run.font.name = "Malgun Gothic"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(*color)


def set_cell(cell, text, size=7.2, bold=False, color=(31, 41, 55), fill=None, align=None):
    if fill:
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(*fill)
    cell.margin_left = Inches(0.025)
    cell.margin_right = Inches(0.025)
    cell.margin_top = Inches(0.012)
    cell.margin_bottom = Inches(0.012)
    frame = cell.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    if align:
        paragraph.alignment = align
    paragraph.font.name = "Malgun Gothic"
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = RGBColor(*color)


def add_compact_slide(prs, case_ids):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = RGBColor(247, 248, 250)

    add_textbox(slide, 0.45, 0.18, 12.4, 0.25, "TradeCare-Agent", size=10, bold=True, color=(75, 94, 115))
    add_textbox(slide, 0.45, 0.48, 12.4, 0.38, "고객 예상 질문 20개 및 예상 답변 핵심", size=22, bold=True, color=(12, 31, 54))
    add_textbox(slide, 0.48, 0.88, 12.2, 0.25, "질문은 발표용으로 핵심 문제명만 압축하고, 답변은 평가에 필요한 Gold Standard 핵심만 표시했습니다.", size=10.2, color=(80, 88, 99))

    rows = 11
    table_shape = slide.shapes.add_table(rows, 6, Inches(0.38), Inches(1.24), Inches(12.6), Inches(5.95))
    table = table_shape.table
    widths = [0.38, 2.12, 3.48, 0.38, 2.12, 3.48]
    for idx, width in enumerate(widths):
        table.columns[idx].width = Inches(width)

    headers = ["No.", "핵심 문제", "예상 답변 핵심", "No.", "핵심 문제", "예상 답변 핵심"]
    for col, header in enumerate(headers):
        set_cell(table.cell(0, col), header, size=7.8, bold=True, color=(255, 255, 255), fill=(29, 78, 116), align=PP_ALIGN.CENTER)

    left = case_ids[:10]
    right = case_ids[10:]
    for row in range(1, 11):
        fill = (255, 255, 255) if row % 2 else (239, 245, 250)
        left_id = left[row - 1]
        right_id = right[row - 1]
        values = [
            left_id,
            COMPACT_QUESTIONS[left_id],
            COMPACT_ANSWERS[left_id],
            right_id,
            COMPACT_QUESTIONS[right_id],
            COMPACT_ANSWERS[right_id],
        ]
        for col, value in enumerate(values):
            set_cell(
                table.cell(row, col),
                value,
                size=7.0 if col in [1, 2, 4, 5] else 7.2,
                bold=col in [0, 3],
                fill=fill,
                align=PP_ALIGN.CENTER if col in [0, 3] else None,
            )


def move_last_slide(prs, index):
    slide_id_list = prs.slides._sldIdLst
    slide_ids = list(slide_id_list)
    last = slide_ids[-1]
    slide_id_list.remove(last)
    slide_id_list.insert(min(index, len(slide_ids)), last)


def main():
    case_ids = load_case_ids()
    for source, target in INPUT_PPTS:
        prs = Presentation(str(source))
        add_compact_slide(prs, case_ids)
        move_last_slide(prs, 11)
        prs.save(str(target))
        print(target)


if __name__ == "__main__":
    main()
